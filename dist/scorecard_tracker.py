VERSION = "1.0.2" 
import pandas as pd
import numpy as np
import tkinter as tk
import tkinter.font as tkFont
from tkinter import ttk, filedialog, messagebox
from tkcalendar import DateEntry
from tksheet import Sheet 
import os
import glob
import json
from datetime import datetime

# Try importing openpyxl for Excel export with colors
try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font
    from openpyxl.utils.dataframe import dataframe_to_rows
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

class ScorecardApp:
    def __init__(self, root):
        self.root = root
        self.root.title(f"DSP Performance Tracker Pro v{VERSION}")
        self.root.geometry("1400x950")

        # Configuration / Settings
        self.settings_file = "settings.json"
        self.settings = {}
        
        # Default metrics to summarize in footer
        self.default_summary_metrics = ["Overall Score", "Packages Delivered", "FICO Metric", "POD", "DCR"]
        
        self.load_settings()
        
        self.df = pd.DataFrame()
        self.filtered_df = pd.DataFrame()
        
        # UPDATED METRICS CONFIGURATION based on request
        # "col" = CSV Column Name
        # "op"  = Operator (< means bad if below, > means bad if above)
        # "default" = Default threshold
        self.metrics = {
            "Overall Score":   {"col": "Overall Score", "op": "<", "default": "90.0"},
            "FICO (Raw)":      {"col": "FICO Metric", "op": "<", "default": "850"}, # Was FICO Score
            "POD (Photo)":     {"col": "POD", "op": "<", "default": "100.0"},       # Was POD Score
            "DCR (Comp)":      {"col": "DCR", "op": "<", "default": "100.0"},       # Was DCR Score
            "CDF (Cust)":      {"col": "CDF DPMO", "op": ">", "default": "0.0"},    # Was CDF DPMO Score
            "DSB (DPMO)":      {"col": "DSB DPMO", "op": ">", "default": "0.0"},    # New
            "CED":             {"col": "CED", "op": ">", "default": "0.0"},         # New
            "Speeding":        {"col": "Speeding Event Rate (per trip)", "op": ">", "default": "0.0"},
            "Seatbelt":        {"col": "Seatbelt-Off Rate (per trip)", "op": ">", "default": "0.0"},
            "Distractions":    {"col": "Distractions Rate (per trip)", "op": ">", "default": "0.0"},
            "Signal Vio":      {"col": "Sign/ Signal Violations Rate (per trip)", "op": ">", "default": "0.0"},
            "Following Dist":  {"col": "Following Distance Rate (per trip)", "op": ">", "default": "0.0"},
        }
        
        # Columns to explicitly HIDE from Viewport
        self.excluded_columns = [
            'FICO Score', 'FICO Tier', 'FICO Metric Weight Applied',
            'Speeding Event Rate Score', 'Speeding Event Rate Tier', 'Speeding Event Rate Weight Applied',
            'Seatbelt-Off Rate Score', 'Seatbelt-Off Rate Tier', 'Seatbelt-Off Rate Weight Applied',
            'Distractions Rate Score', 'Distractions Rate Tier', 'Distractions Rate Weight Applied',
            'Sign/ Signal Violations Rate Score', 'Sign/ Signal Violations Rate Tier', 'Sign/ Signal Violations Rate Weight Applied',
            'Following Distance Rate Score', 'Following Distance Rate Tier', 'Following Distance Rate Weight Applied',
            'CDF DPMO Score', 'CDF DPMO Tier', 'CDF DPMO Weight Applied',
            'CED Score', 'CED Tier', 'CED Weight Applied',
            'DCR Score', 'DCR Tier', 'DCR Weight Applied',
            'DSB DPMO Score', 'DSB DPMO Tier', 'DSB DPMO Weight Applied',
            'POD Score', 'POD Tier', 'POD Weight Applied',
            "PSB", "PSB Score", "PSB Tier", "PSB Weight Applied",
            'Transporter ID', 'Overall Standing'
        ]

        self.threshold_inputs = {} # Stores checkbox state and entry widgets
        self.summary_check_vars = {} # Stores settings checkboxes
        
        self.style = ttk.Style()
        self.setup_ui()
        self.apply_theme()
        self.load_active_data()

    def load_settings(self):
        base_path = os.path.dirname(os.path.abspath(__file__))
        
        defaults = {
            "weekly_path": os.path.join(base_path, "Weekly_Scorecards"),
            "trailing_path": os.path.join(base_path, "Six_Week_Trailing"),
            "summary_metrics": self.default_summary_metrics,
            "theme": "light"
        }

        if os.path.exists(self.settings_file):
            try:
                with open(self.settings_file, 'r') as f:
                    loaded = json.load(f)
                    self.settings = {**defaults, **loaded}
            except Exception as e:
                print(f"Error loading settings: {e}")
                self.settings = defaults
        else:
            self.settings = defaults
            self.save_settings()

    def save_settings(self):
        try:
            with open(self.settings_file, 'w') as f:
                json.dump(self.settings, f, indent=4)
        except Exception as e:
            messagebox.showerror("Error", f"Could not save settings: {e}")

    def apply_theme(self):
        """Applies Light or Dark theme to the UI."""
        theme = self.settings.get("theme", "light")
        
        if theme == "dark":
            # --- PROFESSIONAL DARK THEME ---
            bg_color = "#2e2e2e"       # Soft Dark Gray
            fg_color = "#e0e0e0"       # Off-White (Less harsh than pure white)
            entry_bg = "#454545"       # Lighter Gray for inputs
            header_bg = "#3a3a3a"      # Slightly lighter for headers/buttons
            select_bg = "#007acc"      # VS Code Blue for selection
            sheet_theme = "dark blue"  # Keeps the sheet matching
            
            self.root.configure(bg=bg_color)
            self.style.theme_use('clam') 
            
            # General Settings
            self.style.configure(".", background=bg_color, foreground=fg_color, fieldbackground=entry_bg)
            
            # Buttons - Flatter and Softer
            self.style.configure("TButton", background=header_bg, foreground=fg_color, borderwidth=1, relief="flat")
            self.style.map("TButton", background=[("active", "#505050"), ("pressed", "#202020")], relief=[("pressed", "sunken")])
            
            # LabelFrames - Remove the boxy "Groove" border for a cleaner look
            self.style.configure("TLabelframe", background=bg_color, foreground=fg_color, bordercolor="#454545", relief="solid", borderwidth=1)
            self.style.configure("TLabelframe.Label", background=bg_color, foreground="#00aaff", font=('Arial', 9, 'bold'))
            
            # Inputs
            self.style.configure("TEntry", fieldbackground=entry_bg, foreground="#ffffff", borderwidth=0, padding=5)
            self.style.configure("TCombobox", fieldbackground=entry_bg, foreground="#ffffff", background=header_bg, arrowcolor="white")
            self.style.map("TCombobox", fieldbackground=[("readonly", entry_bg)], selectbackground=[("readonly", select_bg)])
            
            # Checkboxes/Radios
            self.style.configure("TCheckbutton", background=bg_color, foreground=fg_color, indicatorbackground=entry_bg, indicatorcolor=fg_color)
            self.style.configure("TRadiobutton", background=bg_color, foreground=fg_color, indicatorbackground=entry_bg, indicatorcolor=fg_color)
            self.style.map("TCheckbutton", indicatorbackground=[("selected", select_bg)])
            
            # Notebook (Tabs)
            self.style.configure("TNotebook", background=bg_color, borderwidth=0)
            self.style.configure("TNotebook.Tab", background=header_bg, foreground=fg_color, borderwidth=0, padding=[10, 5])
            self.style.map("TNotebook.Tab", background=[("selected", select_bg)], foreground=[("selected", "white")])
            
            # DateEntry Styling (Specific colors to fix visibility)
            style_date = {'background': header_bg, 'foreground': 'white', 'headersbackground': select_bg, 'headersforeground': 'white'}
            self.cal_start.config(**style_date)
            self.cal_end.config(**style_date)

        else:
            # Light Mode Palette (Default)
            self.style.theme_use('vista' if 'vista' in self.style.theme_names() else 'clam')
            sheet_theme = "light green"
            
            # Reset standard colors
            default_bg = "#f0f0f0"
            self.root.configure(bg=default_bg)
            self.style.configure(".", background=default_bg, foreground="black", fieldbackground="white")
            self.style.configure("TLabel", background=default_bg, foreground="black")
            self.style.configure("TLabelframe", background=default_bg, foreground="black")
            self.style.configure("TLabelframe.Label", background=default_bg, foreground="black")
            self.style.configure("TCheckbutton", background=default_bg, foreground="black")
            self.style.configure("TRadiobutton", background=default_bg, foreground="black")
            self.style.configure("TButton", background="#e1e1e1", foreground="black")
            
            # DateEntry Defaults
            self.cal_start.config(background='darkblue', foreground='white')
            self.cal_end.config(background='darkblue', foreground='white')

        # Apply Sheet Theme
        if hasattr(self, 'sheet'):
            self.sheet.change_theme(sheet_theme)
            self.sheet.redraw()

    def setup_ui(self):
        # --- 1. SETTINGS & MODE ---
        data_source_frame = ttk.LabelFrame(self.root, text="Select Data Source")
        data_source_frame.pack(fill="x", padx=10, pady=5)
        data_source_frame.columnconfigure(2, weight=1)

        self.source_var = tk.StringVar(value="weekly")
        ttk.Radiobutton(data_source_frame, text="Weekly Report", variable=self.source_var, value="weekly", command=self.load_active_data).grid(row=0, column=0, padx=10)
        ttk.Radiobutton(data_source_frame, text="6-Week Trailing", variable=self.source_var, value="trailing", command=self.load_active_data).grid(row=0, column=1, padx=10)
        ttk.Button(data_source_frame, text="Settings ⚙️", command=self.open_settings_menu).grid(row=0, column=3, padx=10, pady=8, sticky="e")

        # --- 2. FILTERS ---
        filter_frame = ttk.Frame(self.root)
        filter_frame.pack(fill="x", padx=10, pady=5)

        # Basic Filters
        basic_box = ttk.LabelFrame(filter_frame, text="General Filters")
        basic_box.pack(side="left", fill="y", padx=5)

        ttk.Label(basic_box, text="Employee:").grid(row=0, column=0, padx=5, pady=5)
        self.employee_var = tk.StringVar()
        self.employee_dropdown = ttk.Combobox(basic_box, textvariable=self.employee_var, width=25)
        self.employee_dropdown.grid(row=0, column=1, padx=5)
        # Bind selection change to recalculate summary immediately if desired, or wait for Run Query
        
        ttk.Label(basic_box, text="Start Date:").grid(row=1, column=0, padx=5)
        self.cal_start = DateEntry(basic_box, width=12, background='darkblue', foreground='white', borderwidth=2)
        self.cal_start.grid(row=1, column=1, sticky="w", padx=5, pady=5)

        ttk.Label(basic_box, text="End Date:").grid(row=2, column=0, padx=5)
        self.cal_end = DateEntry(basic_box, width=12, background='darkblue', foreground='white', borderwidth=2)
        self.cal_end.grid(row=2, column=1, sticky="w", padx=5, pady=5)

        # Advanced Thresholds
        adv_box = ttk.LabelFrame(filter_frame, text="Advanced Search (Rows with Violations)")
        adv_box.pack(side="left", fill="both", expand=True, padx=5)

        row_c, col_c = 0, 0
        for label, config in self.metrics.items():
            chk_var = tk.BooleanVar()
            m_frame = ttk.Frame(adv_box)
            m_frame.grid(row=row_c, column=col_c, sticky="w", padx=5, pady=2)
            
            ttk.Checkbutton(m_frame, text=label, variable=chk_var).pack(side="left")
            comp_text = " < " if config["op"] == "<" else " > "
            ttk.Label(m_frame, text=comp_text).pack(side="left")
            ent = ttk.Entry(m_frame, width=6)
            ent.insert(0, config["default"])
            ent.pack(side="left")
            
            self.threshold_inputs[label] = (chk_var, ent)

            col_c += 1
            if col_c > 3: # 4 columns wide
                col_c = 0
                row_c += 1

        # Buttons
        btn_frame = ttk.Frame(filter_frame)
        btn_frame.pack(side="right", fill="y")
        ttk.Button(btn_frame, text="RUN QUERY", command=self.apply_filters).pack(fill="x", pady=8)
        ttk.Button(btn_frame, text="EXPORT EXCEL", command=lambda: self.export_data(mode="standard")).pack(fill="x", pady=8)
        ttk.Button(btn_frame, text="EXPORT BY EMP", command=lambda: self.export_data(mode="employee")).pack(fill="x", pady=8)

        # --- 3. VIEWPORT ---
        view_frame = ttk.Frame(self.root)
        view_frame.pack(fill="both", expand=True, padx=10, pady=5)

        self.sheet = Sheet(view_frame, headers=[], data=[[]], theme="light green", empty_horizontal=0, empty_vertical=0)
        self.sheet.enable_bindings(("single_select", "row_select", "column_width_resize", "arrowkeys", 
                                    "copy", "cut", "paste", "right_click_popup_menu"))
        self.sheet.pack(fill="both", expand=True)

        # --- 4. SUMMARY ---
        self.summary_frame = ttk.LabelFrame(self.root, text="Totals & Averages (Selection)")
        self.summary_frame.pack(fill="x", padx=10, pady=10)
        self.summary_label = ttk.Label(self.summary_frame, text="Ready.", font=('Arial', 10, 'bold'))
        self.summary_label.pack(pady=10)

    def open_settings_menu(self):
        top = tk.Toplevel(self.root)
        top.title("Settings")
        top.geometry("600x450")

        # Apply current theme to popup background
        if self.settings.get("theme") == "dark":
            top.configure(bg="#2b2b2b")

        notebook = ttk.Notebook(top)
        notebook.pack(fill='both', expand=True, padx=5, pady=5)

        # TAB 1: Folders
        tab1 = ttk.Frame(notebook)
        notebook.add(tab1, text='Folders')

        ttk.Label(tab1, text="Weekly Scorecards Folder:").grid(row=0, column=0, padx=10, pady=10, sticky="w")
        entry_weekly = ttk.Entry(tab1, width=40)
        entry_weekly.insert(0, self.settings["weekly_path"])
        entry_weekly.grid(row=0, column=1, padx=10)
        ttk.Button(tab1, text="Browse", command=lambda: self.browse_path(entry_weekly)).grid(row=0, column=2)

        ttk.Label(tab1, text="6-Week Trailing Folder:").grid(row=1, column=0, padx=10, pady=10, sticky="w")
        entry_trailing = ttk.Entry(tab1, width=40)
        entry_trailing.insert(0, self.settings["trailing_path"])
        entry_trailing.grid(row=1, column=1, padx=10)
        ttk.Button(tab1, text="Browse", command=lambda: self.browse_path(entry_trailing)).grid(row=1, column=2)

        # TAB 2: Summary Stats
        tab2 = ttk.Frame(notebook)
        notebook.add(tab2, text='Summary Display')
        
        ttk.Label(tab2, text="Select metrics to show in the footer summary:").pack(pady=5)
        
        # Frame for checkboxes
        chk_frame = ttk.Frame(tab2)
        chk_frame.pack(fill='both', expand=True, padx=10)
        
        # Clear old vars
        self.temp_summary_vars = {}
        
        # Always show mandatory ones
        ttk.Label(chk_frame, text="Always Shown: Overall Score, Packages Delivered").grid(row=0, column=0, columnspan=2, sticky="w", pady=5)
        
        r, c = 1, 0
        possible_metrics = sorted([m["col"] for m in self.metrics.values()])
        for metric in possible_metrics:
            var = tk.BooleanVar(value=(metric in self.settings.get("summary_metrics", [])))
            self.temp_summary_vars[metric] = var
            ttk.Checkbutton(chk_frame, text=metric, variable=var).grid(row=r, column=c, sticky="w", padx=5)
            c += 1
            if c > 1:
                c = 0
                r += 1

        # TAB 3: Appearance
        tab3 = ttk.Frame(notebook)
        notebook.add(tab3, text='Appearance')

        ttk.Label(tab3, text="Application Theme:").grid(row=0, column=0, padx=20, pady=20, sticky="w")
        
        theme_var = tk.StringVar(value=self.settings.get("theme", "light"))
        
        rb_light = ttk.Radiobutton(tab3, text="Light Mode (Default)", variable=theme_var, value="light")
        rb_light.grid(row=1, column=0, padx=20, pady=5, sticky="w")
        
        rb_dark = ttk.Radiobutton(tab3, text="Dark Mode (Professional)", variable=theme_var, value="dark")
        rb_dark.grid(row=2, column=0, padx=20, pady=5, sticky="w")

        # SAVE BUTTON
        def save_settings_btn():
            self.settings["weekly_path"] = entry_weekly.get()
            self.settings["trailing_path"] = entry_trailing.get()

            # Save Theme
            self.settings["theme"] = theme_var.get()

            # Save summary selection
            selected = [m for m, v in self.temp_summary_vars.items() if v.get()]
            self.settings["summary_metrics"] = selected
            
            self.save_settings()
            self.apply_theme()
            self.load_active_data()
            top.destroy()
            messagebox.showinfo("Settings", "Settings saved and theme updated.")

        ttk.Button(top, text="Save & Close", command=save_settings_btn).pack(pady=10)

    def browse_path(self, entry_widget):
        path = filedialog.askdirectory()
        if path:
            entry_widget.delete(0, tk.END)
            entry_widget.insert(0, path)

    def date_to_iso_week(self, dt):
        year, week, day = dt.isocalendar()
        return f"{year}-W{week:02d}"

    def clean_numeric(self, val):
        if pd.isna(val) or val == "": return np.nan
        if isinstance(val, str):
            val = val.replace('%', '').replace(',', '')
        try:
            return float(val)
        except:
            return np.nan

    def load_active_data(self):
        mode = self.source_var.get()
        path = self.settings.get(f"{mode}_path", "")
        
        if not path or not os.path.exists(path):
            self.summary_label.config(text=f"Error: Folder not found for '{mode}'.")
            self.df = pd.DataFrame()
            self.update_viewport()
            return

        files = glob.glob(os.path.join(path, "*.csv"))
        all_dfs = []
        for f in files:
            try:
                temp_df = pd.read_csv(f, encoding='utf-8-sig')
                temp_df.columns = temp_df.columns.str.strip()
                all_dfs.append(temp_df)
            except Exception as e:
                print(f"Error reading {f}: {e}")

        if all_dfs:
            self.df = pd.concat(all_dfs, ignore_index=True)
            
            # Pre-clean numeric columns
            for config in self.metrics.values():
                col_name = config["col"]
                if col_name in self.df.columns:
                    self.df[col_name] = self.df[col_name].apply(self.clean_numeric)

            # Package count cleaning
            if "Packages Delivered" in self.df.columns:
                self.df["Packages Delivered"] = self.df["Packages Delivered"].apply(self.clean_numeric)

            # Update Employee dropdown
            if "Delivery Associate" in self.df.columns:
                employees = sorted(self.df["Delivery Associate"].astype(str).unique().tolist())
                self.employee_dropdown['values'] = ["ALL"] + employees
                self.employee_dropdown.set("ALL")
            
            self.summary_label.config(text=f"Loaded {len(self.df)} rows from {mode}.")
        else:
            self.df = pd.DataFrame()
            self.summary_label.config(text=f"No CSV files found in {path}")
        
        self.filtered_df = pd.DataFrame()
        self.update_viewport()

    def apply_filters(self):
        if self.df.empty: return

        query_df = self.df.copy()

        # 1. Date Filter
        if "Week" in query_df.columns:
            try:
                start_w = self.date_to_iso_week(self.cal_start.get_date())
                end_w = self.date_to_iso_week(self.cal_end.get_date())
                query_df = query_df[(query_df["Week"] >= start_w) & (query_df["Week"] <= end_w)]
            except: pass

        # 2. Employee Filter
        emp = self.employee_var.get()
        if emp != "ALL" and "Delivery Associate" in query_df.columns:
            query_df = query_df[query_df["Delivery Associate"] == emp]

        # 3. Sort: Week Ascending, Overall Score Descending
        if "Week" in query_df.columns and "Overall Score" in query_df.columns:
            query_df = query_df.sort_values(by=["Week", "Overall Score"], ascending=[True, False])

        self.filtered_df = query_df
        self.update_viewport()
        self.calculate_summary()

    def get_active_thresholds(self):
        """Helper to get current limits and operators"""
        active = {}
        for label, (chk_var, ent_widget) in self.threshold_inputs.items():
            config = self.metrics[label]
            try:
                limit = float(ent_widget.get())
            except:
                limit = float(config["default"])
            
            active[config["col"]] = {
                "op": config["op"],
                "limit": limit,
                "checked": chk_var.get() # True if filtering on this
            }
        return active

    def update_viewport(self):
        self.sheet.set_sheet_data([[]])
        self.sheet.dehighlight_all()
        
        if self.filtered_df.empty: return

        # Columns to View
        all_cols = list(self.filtered_df.columns)
        view_cols = [c for c in all_cols if c not in self.excluded_columns]
        
        # Ensure ID columns first
        for col in ["Delivery Associate", "Week"]:
            if col in view_cols:
                view_cols.remove(col)
                view_cols.insert(0, col)

        thresholds = self.get_active_thresholds()
        filter_rows_active = any(t["checked"] for t in thresholds.values())

        final_rows = []
        original_indices = []

        # Iterate rows to determine visibility
        for idx, row in self.filtered_df.iterrows():
            keep_row = True
            
            # Logic: If advanced checkboxes are checked, show row ONLY if it violates at least one checked metric
            if filter_rows_active:
                has_violation = False
                for col_name, rules in thresholds.items():
                    if rules["checked"] and col_name in row:
                        val = row[col_name]
                        if pd.isna(val): continue
                        if (rules["op"] == "<" and val < rules["limit"]) or \
                            (rules["op"] == ">" and val > rules["limit"]):
                            has_violation = True
                            break
                if not has_violation:
                    keep_row = False
            
            if keep_row:
                # Replace nan with blank for display
                display_data = ["" if pd.isna(row[c]) else row[c] for c in view_cols]
                final_rows.append(display_data)
                original_indices.append(idx)

        # Set Data
        self.sheet.headers(view_cols)
        self.sheet.set_sheet_data(final_rows)

        # Apply Highlight colors
        bg_color = "#ffcccc" # Light Red
        if self.settings.get("theme") == "dark":
            bg_color = "#8b3a3a"  # Darker red for dark mode to be readable
        
        for r_idx, row_idx in enumerate(original_indices):
            row_data = self.filtered_df.loc[row_idx]
            for c_idx, col_name in enumerate(view_cols):
                if col_name in thresholds:
                    rules = thresholds[col_name]
                    val = row_data[col_name]
                    
                    if not pd.isna(val):
                        if (rules["op"] == "<" and val < rules["limit"]) or \
                            (rules["op"] == ">" and val > rules["limit"]):
                            self.sheet.highlight_cells(row=r_idx, column=c_idx, bg=bg_color)

        self.summary_label.config(text=f"Rows Displayed: {len(final_rows)}")
        self.sheet.set_all_cell_sizes_to_text()

    def calculate_summary(self):
        if self.filtered_df.empty:
            self.summary_label.config(text="No data selected.")
            return

        # Mandatory Metrics
        txt_parts = [f"Count: {len(self.filtered_df)}"]
        
        if "Packages Delivered" in self.filtered_df.columns:
            total_pkgs = self.filtered_df["Packages Delivered"].sum()
            txt_parts.append(f"Total Pkgs: {int(total_pkgs):,}")
            
        if "Overall Score" in self.filtered_df.columns:
            avg_overall = self.filtered_df["Overall Score"].mean()
            txt_parts.append(f"Overall Score: {avg_overall:.2f}")

        # Configured Metrics from Settings
        user_metrics = self.settings.get("summary_metrics", [])
        for col in user_metrics:
            if col in self.filtered_df.columns:
                # Special logic: "Rate" usually average, "DPMO" usually average
                # Just doing average for all scores/rates
                vals = self.filtered_df[col].dropna()
                if not vals.empty:
                    avg_val = vals.mean()
                    txt_parts.append(f"{col}: {avg_val:.2f}")

        self.summary_label.config(text=" | ".join(txt_parts))

    def export_data(self, mode="standard"):
        if self.filtered_df.empty:
            messagebox.showwarning("Export", "No data to export.")
            return
            
        if not HAS_OPENPYXL:
            messagebox.showerror("Error", "Library 'openpyxl' not found.\nPlease install it to export with formatting:\npip install openpyxl")
            return

        fp = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel Files", "*.xlsx")])
        if not fp: return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Scorecard Data"
            
            red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
            bold_font = Font(bold=True)
            
            thresholds = self.get_active_thresholds()
            
            # --- PREPARE DATA ---
            view_cols = self.sheet.headers() # Get columns currently in viewport
            
            # Helper to check violations for coloring
            def is_violation(col, val):
                if col in thresholds and val is not None and val != "":
                    rules = thresholds[col]
                    try:
                        f_val = float(val)
                        if rules["op"] == "<" and f_val < rules["limit"]: return True
                        if rules["op"] == ">" and f_val > rules["limit"]: return True
                    except: pass
                return False

            if mode == "standard":
                # Header
                ws.append(view_cols)
                for cell in ws[1]: cell.font = bold_font
                
                # Rows (using the Sheet data directly to match sort/filter)
                sheet_data = self.sheet.get_sheet_data()
                
                for r_idx, row_data in enumerate(sheet_data, start=2):
                    for c_idx, val in enumerate(row_data):
                        # Convert empty strings back to None for Excel or keep as text
                        cell = ws.cell(row=r_idx, column=c_idx+1, value=val)
                        col_name = view_cols[c_idx]
                        if is_violation(col_name, val):
                            cell.fill = red_fill

            elif mode == "employee":
                # Group by Employee
                if "Delivery Associate" not in view_cols:
                    messagebox.showerror("Error", "Delivery Associate column missing.")
                    return

                employees = self.filtered_df["Delivery Associate"].unique()
                
                current_row = 1
                # Header
                for c_idx, col in enumerate(view_cols):
                    cell = ws.cell(row=current_row, column=c_idx+1, value=col)
                    cell.font = bold_font
                current_row += 1
                
                for emp in employees:
                    emp_data = self.filtered_df[self.filtered_df["Delivery Associate"] == emp].copy()
                    # Ensure sorted by Week
                    if "Week" in emp_data.columns:
                        emp_data = emp_data.sort_values(by="Week")
                    
                    # Write Employee Rows
                    # Need to map emp_data to view_cols
                    start_row_group = current_row
                    for _, row in emp_data.iterrows():
                        for c_idx, col in enumerate(view_cols):
                            val = row[col] if col in row else ""
                            if pd.isna(val): val = ""
                            
                            cell = ws.cell(row=current_row, column=c_idx+1, value=val)
                            
                            # Highlight logic
                            if is_violation(col, val):
                                cell.fill = red_fill
                        current_row += 1
                    
                    # Totals Row
                    ws.cell(row=current_row, column=1, value=f"{emp} Totals/Avgs").font = bold_font
                    
                    for c_idx, col in enumerate(view_cols):
                        if col == "Delivery Associate" or col == "Week": continue
                        
                        # Calculate avg or sum
                        if col == "Packages Delivered":
                            val = emp_data[col].sum()
                        elif col in self.metrics:
                            val = emp_data[col].mean()
                        else:
                            val = ""
                            
                        if isinstance(val, (int, float)):
                            ws.cell(row=current_row, column=c_idx+1, value=round(val, 2)).font = bold_font
                            
                    current_row += 2 # Add space between employees

            # Auto-adjust column width
            for column_cells in ws.columns:
                length = max(len(str(cell.value) or "") for cell in column_cells)
                ws.column_dimensions[column_cells[0].column_letter].width = length + 2

            wb.save(fp)
            messagebox.showinfo("Success", f"Exported to {fp}")

        except Exception as e:
            messagebox.showerror("Export Error", str(e))

def main():
    root = tk.Tk()
    app = ScorecardApp(root)
    root.mainloop()

if __name__ == "__main__":
    main()